from django.contrib import messages
from django.contrib.auth import authenticate, get_user_model, login
from django.contrib.auth.models import Group
from django.db import transaction

from django.db.models import Count, F, Q, ExpressionWrapper, IntegerField
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl.styles import Alignment

from .decorators import permission_required
from .models import ActivityLog
from blog.models import *
import jdatetime

from accounts.forms import *
from openpyxl import Workbook
from django.http import HttpResponse
from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.template.loader import get_template
from decimal import Decimal, InvalidOperation

User = get_user_model()


# =========================================================
# ROLE CONSTANTS
# =========================================================

SYSTEM_MANAGER = "مدیر سیستم"
CONTENT_MANAGER = "مدیر محتوا"
USER_MANAGER = "مدیر کاربران"

ROLE_CHOICES = [
    ("system", SYSTEM_MANAGER),
    ("content", CONTENT_MANAGER),
    ("users", USER_MANAGER),
]


# =========================================================
# ROLE HELPERS
# =========================================================

def get_user_roles(user):
    """
    نقش‌های فعلی کاربر را برمی‌گرداند.
    یک کاربر می‌تواند چند نقش داشته باشد.
    """

    if user.is_superuser:
        roles = {SYSTEM_MANAGER}
    else:
        roles = set(
            user.groups.values_list(
                "name",
                flat=True
            )
        )

    return roles


def set_user_roles(user, selected_roles):
    """
    نقش‌های کاربر را بر اساس انتخاب‌های جدید تنظیم می‌کند.

    selected_roles:
        ["system", "content", "users"]
    """

    selected_roles = set(selected_roles)

    system_selected = "system" in selected_roles
    content_selected = "content" in selected_roles
    users_selected = "users" in selected_roles

    # -----------------------------------------
    # مدیر سیستم
    # -----------------------------------------

    if system_selected:
        system_group, _ = Group.objects.get_or_create(
            name=SYSTEM_MANAGER
        )

        user.groups.clear()
        user.groups.add(system_group)

        user.is_staff = True
        user.is_superuser = True

    # -----------------------------------------
    # سایر نقش‌ها
    # -----------------------------------------

    else:
        groups = []

        if content_selected:
            content_group, _ = Group.objects.get_or_create(
                name=CONTENT_MANAGER
            )
            groups.append(content_group)

        if users_selected:
            users_group, _ = Group.objects.get_or_create(
                name=USER_MANAGER
            )
            groups.append(users_group)

        user.groups.set(groups)

        user.is_superuser = False
        user.is_staff = bool(groups)

    user.save(
        update_fields=[
            "is_staff",
            "is_superuser",
        ]
    )


# =========================================================
# DASHBOARD
# =========================================================

@permission_required("accounts.access_dashboard")
def dashboard_home(request):
    books_count = Book.objects.count()

    categories_count = Subject.objects.count()

    available_count = (
        Book.objects
        .annotate(
            borrowed_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
            )
        )
        .filter(
            total_copies__gt=F("borrowed_count")
        )
        .count()
    )

    unavailable_count = (
        Book.objects
        .annotate(
            borrowed_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
            )
        )
        .filter(
            total_copies__lte=F("borrowed_count")
        )
        .count()
    )

    active_loans_count = Loan.objects.filter(
        status=Loan.STATUS_BORROWED
    ).count()

    pending_loans_count = 0

    latest_books = (
        Book.objects
        .select_related(
            "subject",
            "publisher",
            "location",
        )
        .order_by("-created_at")[:5]
    )

    context = {
        "books_count": books_count,
        "categories_count": categories_count,
        "available_count": available_count,
        "unavailable_count": unavailable_count,
        "active_loans_count": active_loans_count,
        "pending_loans_count": pending_loans_count,
        "latest_books": latest_books,
        "can_manage_books": request.user.has_perm(
            "blog.manage_books"
        ),
        "can_manage_categories": request.user.has_perm(
            "blog.manage_categories"
        ),
        "can_manage_loans": request.user.has_perm(
            "blog.manage_loans"
        ),
        "can_manage_users": request.user.has_perm(
            "accounts.change_user"
        ),
    }

    return render(
        request,
        "dashboard/index.html",
        context
    )


# =========================================================
# BOOKS
# =========================================================

def normalize_digits(value):
    """
    تبدیل اعداد فارسی و عربی به اعداد انگلیسی.
    """

    if not value:
        return value

    translation_table = str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789"
    )

    return value.translate(translation_table)



def get_book_form_context(book=None, form_data=None):
    context = {
        "book": book,
        "edit_mode": book is not None,
        "subjects": Subject.objects.all().order_by("name"),
        "categories": Category.objects.all().order_by("name"),
        "authors": Author.objects.all().order_by("name"),
        "translators": Translator.objects.all().order_by("name"),
        "publishers": Publisher.objects.all().order_by("name"),
        "locations": BookLocation.objects.all().order_by("code"),
    }

    if form_data:
        context["form_data"] = form_data
        context["author_values"] = form_data.getlist("authors")
        context["translator_values"] = form_data.getlist("translators")

    return context




def get_filtered_books(request):
    query = request.GET.get("q", "").strip()
    search_by = request.GET.get("search_by", "title").strip()
    sort = request.GET.get("sort", "newest").strip()
    availability = request.GET.get("availability", "").strip()

    valid_search_fields = {
        "title",
        "isbn",
        "registration_number",
        "author",
        "translator",
        "publisher",
        "subject",
        "category",
        "location",
        "description",
    }

    if search_by not in valid_search_fields:
        search_by = "title"

    valid_sort_options = {
        "newest",
        "oldest",
        "title_asc",
        "title_desc",
    }

    if sort not in valid_sort_options:
        sort = "newest"

    books = (
        Book.objects
        .select_related(
            "subject",
            "publisher",
            "location",
            "category",
        )
        .prefetch_related(
            "authors",
            "translators",
        )
        .annotate(
            borrowed_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
                distinct=True,
            )
        )
        .annotate(
            available_copies=ExpressionWrapper(
                F("total_copies") - F("borrowed_count"),
                output_field=IntegerField(),
            )
        )
    )

    if query:
        if search_by == "title":
            books = books.filter(title__icontains=query)

        elif search_by == "isbn":
            books = books.filter(isbn__icontains=query)

        elif search_by == "registration_number":
            normalized_query = normalize_digits(query)
            if normalized_query.isdigit():
                books = books.filter(
                    registration_number=int(normalized_query)
                )
            else:
                books = books.none()

        elif search_by == "author":
            books = books.filter(
                authors__name__icontains=query
            )

        elif search_by == "translator":
            books = books.filter(
                translators__name__icontains=query
            )

        elif search_by == "publisher":
            books = books.filter(
                publisher__name__icontains=query
            )

        elif search_by == "subject":
            books = books.filter(
                subject__name__icontains=query
            )

        elif search_by == "category":
            books = books.filter(
                category__name__icontains=query
            )

        elif search_by == "location":
            books = books.filter(
                Q(location__code__icontains=query)
                | Q(location__description__icontains=query)
            )

        elif search_by == "description":
            books = books.filter(
                description__icontains=query
            )

        books = books.distinct()

    if sort == "newest":
        books = books.order_by("-created_at")

    elif sort == "oldest":
        books = books.order_by("created_at")

    elif sort == "title_asc":
        books = books.order_by("title")

    elif sort == "title_desc":
        books = books.order_by("-title")

    if availability == "available":
        books = books.filter(
            available_copies__gt=0
        )

    elif availability == "unavailable":
        books = books.filter(
            available_copies__lte=0
        )

    return books, {
        "query": query,
        "search_by": search_by,
        "sort": sort,
        "availability": availability,
    }




@permission_required("blog.manage_books")
def manage_books(request):
    books, filters = get_filtered_books(request)

    context = {
        "books": books,
        **filters,
    }

    return render(
        request,
        "dashboard/books.html",
        context
    )




@permission_required("blog.manage_books")
def export_books_excel(request):
    books, filters = get_filtered_books(request)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "کتاب‌ها"

    # راست‌چین کردن شیت
    worksheet.sheet_view.rightToLeft = True

    headers = [
        "شماره ثبت",
        "عنوان کتاب",
        "شابک",
        "مؤلفان",
        "مترجمان",
        "موضوع",
        "ناشر",
        "سال انتشار",
        "محل قرارگیری",
        "تعداد نسخه",
        "نسخه‌های امانت",
        "نسخه‌های موجود",
    ]

    worksheet.append(headers)

    for book in books:
        authors = "، ".join(
            author.name
            for author in book.authors.all()
        )

        translators = "، ".join(
            translator.name
            for translator in book.translators.all()
        )

        subject = (
            book.subject.name
            if book.subject
            else ""
        )

        publisher = (
            book.publisher.name
            if book.publisher
            else ""
        )

        location = (
            book.location.code
            if book.location
            else ""
        )

        borrowed_count = book.borrowed_count or 0

        available_copies = max(
            book.available_copies or 0,
            0,
        )

        worksheet.append([
            book.registration_number or "",
            book.title or "",
            book.isbn or "",
            authors,
            translators,
            subject,
            publisher,
            book.publication_year or "",
            location,
            book.total_copies or 0,
            borrowed_count,
            available_copies,
        ])

    # -------------------------
    # تنظیم سلول‌ها
    # -------------------------

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
                wrap_text=True,
            )

    # -------------------------
    # عرض ستون‌ها
    # -------------------------

    column_widths = {
        "A": 14,
        "B": 35,
        "C": 20,
        "D": 35,
        "E": 35,
        "F": 25,
        "G": 25,
        "H": 15,
        "I": 18,
        "J": 15,
        "K": 15,
        "L": 15,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # -------------------------
    # تنظیمات نهایی Excel
    # -------------------------

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # ارتفاع ردیف عنوان
    worksheet.row_dimensions[1].height = 30

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="books_list.xlsx"'
    )

    workbook.save(response)

    return response



@permission_required("blog.manage_books")
def print_books(request):
    books, filters = get_filtered_books(request)

    context = {
        "books": books,
        **filters,
    }

    return render(
        request,
        "dashboard/books_print.html",
        context
    )




@permission_required("blog.manage_books")
def add_book(request):
    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        isbn = request.POST.get("isbn", "").strip()
        subject_name = request.POST.get("subject", "").strip()
        category_id = request.POST.get("category", "").strip()
        publisher_name = request.POST.get("publisher", "").strip()
        description = request.POST.get("description", "").strip()
        cover = request.FILES.get("cover")

        # قیمت کتاب؛ اختیاری
        price_raw = normalize_digits(
            request.POST.get("price", "").strip()
        )
        price = None

        total_copies = normalize_digits(
            request.POST.get("total_copies", "1").strip()
        )
        publication_year = normalize_digits(
            request.POST.get("publication_year", "").strip()
        )

        author_names = [
            name.strip()
            for name in request.POST.getlist("authors")
            if name.strip()
        ]

        translator_names = [
            name.strip()
            for name in request.POST.getlist("translators")
            if name.strip()
        ]

        form_data = request.POST

        def render_form_error(message):
            messages.error(request, message)
            return render(
                request,
                "dashboard/book_form.html",
                get_book_form_context(form_data=form_data),
            )

        if not title:
            return render_form_error("عنوان کتاب را وارد کنید.")

        if isbn and Book.objects.filter(isbn=isbn).exists():
            return render_form_error(
                "این شابک قبلاً برای یک کتاب ثبت شده است."
            )

        if not author_names:
            return render_form_error(
                "حداقل یک مؤلف برای کتاب وارد کنید."
            )

        if not subject_name:
            return render_form_error("موضوع کتاب را وارد کنید.")

        category = None
        if category_id:
            category = Category.objects.filter(
                pk=category_id
            ).first()

            if category is None:
                return render_form_error(
                    "دسته‌بندی انتخاب‌شده معتبر نیست. دوباره انتخاب کنید."
                )


        if not publisher_name:
            return render_form_error("نام ناشر را وارد کنید.")

        # اعتبارسنجی قیمت؛ خالی بودن آن مجاز است
        if price_raw:
            try:
                price = Decimal(price_raw)

                if not price.is_finite() or price < 0:
                    return render_form_error(
                        "قیمت کتاب باید عددی معتبر و غیرمنفی باشد."
                    )

            except (InvalidOperation, ValueError, TypeError):
                return render_form_error(
                    "قیمت کتاب باید عددی معتبر و غیرمنفی باشد."
                )

        try:
            total_copies = int(total_copies)
        except (TypeError, ValueError):
            return render_form_error(
                "تعداد نسخه‌ها باید یک عدد باشد."
            )

        if total_copies < 1:
            return render_form_error(
                "تعداد نسخه‌ها باید حداقل ۱ باشد."
            )

        if not publication_year:
            return render_form_error(
                "سال انتشار کتاب را وارد کنید."
            )

        try:
            publication_year = int(publication_year)
        except (TypeError, ValueError):
            return render_form_error(
                "سال انتشار باید یک عدد باشد."
            )

        if not 1000 <= publication_year <= 1600:
            return render_form_error(
                "سال انتشار شمسی واردشده معتبر نیست."
            )




        with transaction.atomic():
            subject, _ = Subject.objects.get_or_create(
                name=subject_name
            )

            publisher, _ = Publisher.objects.get_or_create(
                name=publisher_name
            )

            book = Book.objects.create(
                title=title,
                isbn=isbn or None,
                subject=subject,
                publisher=publisher,
                publication_year=publication_year,
                category=category,
                description=description,
                cover=cover,
                total_copies=total_copies,
                price=price,
                is_available=True,
            )

            for name in author_names:
                author, _ = Author.objects.get_or_create(
                    name=name
                )
                book.authors.add(author)

            for name in translator_names:
                translator, _ = Translator.objects.get_or_create(
                    name=name
                )
                book.translators.add(translator)

        log_activity(
            request=request,
            action=ActivityLog.ACTION_CREATE,
            section="کتاب‌ها",
            description=(
                f"کتاب «{book.title}» با شماره ثبت "
                f"{book.registration_number} اضافه شد."
            ),
            object_id=book.pk,
        )

        messages.success(
            request,
            "کتاب با موفقیت اضافه شد."
        )

        return redirect("manage_books")

    return render(
        request,
        "dashboard/book_form.html",
        get_book_form_context(),
    )






@permission_required("blog.manage_books")
def edit_book(request, pk):
    book = get_object_or_404(
        Book.objects.select_related(
            "subject",
            "publisher",
            "category",
            "location",
        ).prefetch_related(
            "authors",
            "translators",
        ),
        pk=pk,
    )

    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        isbn = request.POST.get("isbn", "").strip()
        subject_name = request.POST.get("subject", "").strip()
        category_id = request.POST.get("category", "").strip()
        publisher_name = request.POST.get("publisher", "").strip()
        description = request.POST.get("description", "").strip()
        cover = request.FILES.get("cover")

        total_copies = normalize_digits(
            request.POST.get(
                "total_copies",
                str(book.total_copies),
            ).strip()
        )

        publication_year = normalize_digits(
            request.POST.get("publication_year", "").strip()
        )

        author_names = [
            name.strip()
            for name in request.POST.getlist("authors")
            if name.strip()
        ]

        translator_names = [
            name.strip()
            for name in request.POST.getlist("translators")
            if name.strip()
        ]

        form_data = request.POST

        def render_form_error(message):
            messages.error(request, message)
            return render(
                request,
                "dashboard/book_form.html",
                get_book_form_context(book, form_data),
            )

        if not title:
            return render_form_error("عنوان کتاب را وارد کنید.")

        if isbn and Book.objects.filter(isbn=isbn).exclude(
            pk=book.pk
        ).exists():
            return render_form_error(
                "این شابک قبلاً برای کتاب دیگری ثبت شده است."
            )

        if not author_names:
            return render_form_error(
                "حداقل یک مؤلف برای کتاب وارد کنید."
            )

        if not subject_name:
            return render_form_error("موضوع کتاب را وارد کنید.")

        category = None
        if category_id:
            category = Category.objects.filter(
                pk=category_id
            ).first()

            if category is None:
                return render_form_error(
                    "دسته‌بندی انتخاب‌شده معتبر نیست. دوباره انتخاب کنید."
                )

        if not publisher_name:
            return render_form_error("نام ناشر را وارد کنید.")

        try:
            total_copies = int(total_copies)
        except (TypeError, ValueError):
            return render_form_error(
                "تعداد نسخه‌ها باید یک عدد باشد."
            )

        if total_copies < 1:
            return render_form_error(
                "تعداد نسخه‌ها باید حداقل ۱ باشد."
            )

        if not publication_year:
            return render_form_error(
                "سال انتشار کتاب را وارد کنید."
            )

        try:
            publication_year = int(publication_year)
        except (TypeError, ValueError):
            return render_form_error(
                "سال انتشار باید یک عدد باشد."
            )

        if not 1000 <= publication_year <= 1600:
            return render_form_error(
                "سال انتشار شمسی واردشده معتبر نیست."
            )



        with transaction.atomic():
            borrowed_count = Loan.objects.filter(
                book=book,
                status=Loan.STATUS_BORROWED,
            ).count()

            if total_copies < borrowed_count:
                return render_form_error(
                    f"تعداد نسخه‌ها نمی‌تواند کمتر از "
                    f"{borrowed_count} نسخه‌ای باشد که "
                    f"در حال حاضر امانت هستند."
                )

            subject, _ = Subject.objects.get_or_create(
                name=subject_name
            )

            publisher, _ = Publisher.objects.get_or_create(
                name=publisher_name
            )

            book.title = title
            book.isbn = isbn or None
            book.subject = subject
            book.publisher = publisher
            book.publication_year = publication_year
            book.category = category
            book.description = description
            book.total_copies = total_copies
            book.is_available = borrowed_count < total_copies

            if cover:
                book.cover = cover

            book.save()

            authors = []

            for name in author_names:
                author, _ = Author.objects.get_or_create(
                    name=name
                )
                authors.append(author)

            book.authors.set(authors)

            translators = []

            for name in translator_names:
                translator, _ = Translator.objects.get_or_create(
                    name=name
                )
                translators.append(translator)

            book.translators.set(translators)

        log_activity(
            request=request,
            action=ActivityLog.ACTION_UPDATE,
            section="کتاب‌ها",
            description=(
                f"اطلاعات کتاب «{book.title}» با شماره ثبت "
                f"{book.registration_number} ویرایش شد."
            ),
            object_id=book.pk,
        )

        messages.success(
            request,
            "اطلاعات کتاب با موفقیت ویرایش شد."
        )

        return redirect("manage_books")

    return render(
        request,
        "dashboard/book_form.html",
        get_book_form_context(book),
    )


@permission_required("blog.manage_books")
def delete_book(request, pk):

    book = get_object_or_404(
        Book,
        pk=pk
    )

    if request.method == "POST":

        book_title = book.title
        book_registration_number = book.registration_number
        book_pk = book.pk

        book.delete()

        log_activity(
            request=request,
            action=ActivityLog.ACTION_DELETE,
            section="کتاب‌ها",
            description=(
                f"کتاب «{book_title}» با شماره ثبت "
                f"{book_registration_number} حذف شد."
            ),
            object_id=book_pk,
        )

        messages.success(
            request,
            "کتاب با موفقیت حذف شد."
        )

        return redirect("manage_books")

    return render(
        request,
        "dashboard/book_delete.html",
        {
            "book": book,
        }
    )


# =========================================================
# CATEGORIES / SUBJECTS
# =========================================================

@permission_required("blog.manage_categories")
def manage_categories(request):
    categories = Category.objects.all().order_by("name")

    return render(
        request,
        "dashboard/categories.html",
        {
            "categories": categories,
        },
    )


@permission_required("blog.manage_categories")
def add_category(request):
    if request.method == "POST":
        name = (
            request.POST.get("name")
            or request.POST.get("category")
            or ""
        ).strip()

        if not name:
            messages.error(
                request,
                "نام دسته‌بندی را وارد کنید."
            )

            return render(
                request,
                "dashboard/category_form.html",
                {
                    "category": None,
                    "edit_mode": False,
                },
            )

        category = Category(name=name)
        category.save()

        log_activity(
            request=request,
            action=ActivityLog.ACTION_CREATE,
            section="دسته‌بندی‌ها",
            description=(
                f"دسته‌بندی «{category.name}» ایجاد شد."
            ),
            object_id=category.pk,
        )

        messages.success(
            request,
            "دسته‌بندی با موفقیت اضافه شد."
        )

        return redirect("manage_categories")

    return render(
        request,
        "dashboard/category_form.html",
        {
            "category": None,
            "edit_mode": False,
        },
    )



@permission_required("blog.manage_categories")
def edit_category(request, pk):
    category = get_object_or_404(Category, pk=pk)

    if request.method == "POST":
        name = (
            request.POST.get("name")
            or request.POST.get("category")
            or ""
        ).strip()

        if not name:
            messages.error(
                request,
                "نام دسته‌بندی را وارد کنید."
            )

            return render(
                request,
                "dashboard/category_form.html",
                {
                    "category": category,
                    "edit_mode": True,
                },
            )

        old_name = category.name

        category.name = name
        category.save()

        log_activity(
            request=request,
            action=ActivityLog.ACTION_UPDATE,
            section="دسته‌بندی‌ها",
            description=(
                f"دسته‌بندی «{old_name}» "
                f"به «{category.name}» تغییر کرد."
            ),
            object_id=category.pk,
        )

        messages.success(
            request,
            "دسته‌بندی با موفقیت ویرایش شد."
        )

        return redirect("manage_categories")

    return render(
        request,
        "dashboard/category_form.html",
        {
            "category": category,
            "edit_mode": True,
        },
    )



@permission_required("blog.manage_categories")
def delete_category(request, pk):
    category = get_object_or_404(Category, pk=pk)

    if request.method == "POST":
        category_name = category.name
        category_id = category.pk

        try:
            category.delete()
        except Exception:
            messages.error(
                request,
                "حذف دسته‌بندی انجام نشد. ممکن است این دسته‌بندی "
                "به اطلاعات دیگری وابسته باشد."
            )
            return redirect("manage_categories")

        log_activity(
            request=request,
            action=ActivityLog.ACTION_DELETE,
            section="دسته‌بندی‌ها",
            description=f"دسته‌بندی «{category_name}» حذف شد.",
            object_id=category_id,
        )

        messages.success(
            request,
            "دسته‌بندی با موفقیت حذف شد."
        )
        return redirect("manage_categories")

    return render(
        request,
        "dashboard/category_delete.html",
        {"category": category},
    )



# =========================================================
# SUPPORT / ACTIVITY LOG HELPER
# =========================================================

def log_activity(
    request,
    action,
    section,
    description,
    object_id="",
):
    """
    ثبت یک رویداد در سیستم پشتیبانی.
    """

    actor = None

    if request.user.is_authenticated:
        actor = request.user

    ip_address = request.META.get(
        "REMOTE_ADDR"
    )

    ActivityLog.objects.create(
        actor=actor,
        action=action,
        section=section,
        description=description,
        object_id=str(object_id) if object_id else "",
        ip_address=ip_address,
    )


# =========================================================
# LOANS
# =========================================================

@permission_required("blog.manage_loans")
def export_loans_excel(request):
    """
    خروجی Excel از تمام سوابق امانت‌ها.
    """
    loans = (
        Loan.objects
        .select_related(
            "user",
            "book",
            "book__subject",
            "book__publisher",
        )
        .prefetch_related(
            "book__authors",
            "book__translators",
        )
        .order_by("-created_at")
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "امانت‌ها"

    headers = [
        "شماره ثبت امانت",
        "کتاب",
        "کاربر",
        "شماره موبایل",
        "وضعیت",
        "تاریخ دریافت",
        "مهلت بازگشت",
        "تاریخ برگشت",
        "تعداد روز دیرکرد",
        "جریمه هر روز",
        "مبلغ جریمه",
        "توضیحات مدیر",
    ]

    worksheet.append(headers)

    def jalali_datetime(value):
        if not value:
            return ""
        try:
            if timezone.is_aware(value):
                value = timezone.localtime(value)
            return jdatetime.datetime.fromgregorian(
                datetime=value
            ).strftime("%Y/%m/%d - %H:%M")
        except Exception:
            return ""

    for loan in loans:
        if loan.status == Loan.STATUS_BORROWED:
            late_days = loan.current_late_days
            current_fine = loan.current_fine
        else:
            late_days = loan.late_days
            current_fine = loan.total_fine

        worksheet.append([
            loan.registration_number,
            loan.book.title if loan.book else "",
            (
                loan.user.get_full_name()
                if loan.user and loan.user.get_full_name()
                else loan.user.username
                if loan.user
                else ""
            ),
            loan.user.phone_number if loan.user else "",
            loan.get_status_display(),
            jalali_datetime(loan.borrowed_at),
            jalali_datetime(loan.due_date),
            jalali_datetime(loan.returned_at),
            late_days,
            loan.daily_fine,
            current_fine,
            loan.admin_note or "",
        ])

    for cell in worksheet[1]:
        cell.alignment = Alignment(
            horizontal="right",
            vertical="center",
            wrap_text=True,
        )
        cell.font = cell.font.copy(bold=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
                wrap_text=True,
            )

    column_widths = {
        "A": 16,
        "B": 35,
        "C": 25,
        "D": 18,
        "E": 18,
        "F": 22,
        "G": 22,
        "H": 22,
        "I": 18,
        "J": 18,
        "K": 18,
        "L": 40,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.rightToLeft = True
    worksheet.row_dimensions[1].height = 30

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="loans_list.xlsx"'
    )

    workbook.save(response)

    return response


@permission_required("blog.manage_loans")
def manage_loans(request):

    available_books = (
        Book.objects
        .annotate(
            borrowed_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
            )
        )
        .annotate(
            available_copies=ExpressionWrapper(
                F("total_copies") -
                F("borrowed_count"),
                output_field=IntegerField(),
            )
        )
        .filter(
            available_copies__gt=0
        )
        .select_related(
            "category",
        )
        .order_by(
            "title",
        )
    )

    users = (
        User.objects
        .filter(
            is_active=True,
        )
        .order_by(
            "first_name",
            "last_name",
            "username",
        )
    )

    loans = (
        Loan.objects
        .select_related(
            "user",
            "book",
            "book__category",
        )
        .prefetch_related(
            "book__authors",
        )
        .order_by(
            "-created_at",
        )
    )

    # سال جاری شمسی و چند سال آینده
    today_jalali = jdatetime.date.today()

    current_jalali_year = today_jalali.year

    due_years = range(
        current_jalali_year,
        current_jalali_year + 6,
    )

    context = {
        "available_books": available_books,
        "users": users,
        "loans": loans,
        "due_years": due_years,
    }

    return render(
        request,
        "dashboard/loans.html",
        context,
    )


@permission_required("blog.manage_loans")
def create_loan(request):

    if request.method != "POST":
        return redirect("manage_loans")

    book_id = request.POST.get(
        "book_id",
        "",
    ).strip()

    user_id = request.POST.get(
        "user_id",
        "",
    ).strip()

    due_year = request.POST.get(
        "due_year",
        "",
    ).strip()

    due_month = request.POST.get(
        "due_month",
        "",
    ).strip()

    due_day = request.POST.get(
        "due_day",
        "",
    ).strip()

    due_hour = request.POST.get(
        "due_hour",
        "",
    ).strip()

    due_minute = request.POST.get(
        "due_minute",
        "",
    ).strip()

    daily_fine = request.POST.get(
        "daily_fine",
        "0",
    ).strip()

    admin_note = request.POST.get(
        "admin_note",
        "",
    ).strip()

    # -------------------------
    # بررسی کتاب و کاربر
    # -------------------------

    if not book_id or not user_id:

        messages.error(
            request,
            "ابتدا کتاب و کاربر را انتخاب کنید.",
        )

        return redirect("manage_loans")

    # -------------------------
    # بررسی تاریخ
    # -------------------------

    if not all(
        [
            due_year,
            due_month,
            due_day,
            due_hour,
            due_minute,
        ]
    ):

        messages.error(
            request,
            "تاریخ و ساعت مهلت بازگشت را کامل مشخص کنید.",
        )

        return redirect("manage_loans")

    try:

        due_year = int(due_year)
        due_month = int(due_month)
        due_day = int(due_day)
        due_hour = int(due_hour)
        due_minute = int(due_minute)

    except (TypeError, ValueError):

        messages.error(
            request,
            "تاریخ مهلت بازگشت معتبر نیست.",
        )

        return redirect("manage_loans")

    # -------------------------
    # اعتبارسنجی ساعت
    # -------------------------

    if not 0 <= due_hour <= 23:

        messages.error(
            request,
            "ساعت انتخاب‌شده معتبر نیست.",
        )

        return redirect("manage_loans")

    if not 0 <= due_minute <= 59:

        messages.error(
            request,
            "دقیقه انتخاب‌شده معتبر نیست.",
        )

        return redirect("manage_loans")

    # -------------------------
    # تبدیل تاریخ شمسی به میلادی
    # -------------------------

    try:

        jalali_due_date = jdatetime.datetime(
            due_year,
            due_month,
            due_day,
            due_hour,
            due_minute,
            0,
        )

        gregorian_due_date = (
            jalali_due_date
            .togregorian()
        )

        due_date_value = timezone.make_aware(
            gregorian_due_date,
            timezone.get_current_timezone(),
        )

    except (ValueError, TypeError):

        messages.error(
            request,
            "تاریخ انتخاب‌شده معتبر نیست.",
        )

        return redirect("manage_loans")

    # -------------------------
    # مهلت نباید در گذشته باشد
    # -------------------------

    now = timezone.now()

    if due_date_value <= now:

        messages.error(
            request,
            "مهلت بازگشت باید بعد از زمان فعلی باشد.",
        )

        return redirect("manage_loans")

    # -------------------------
    # جریمه روزانه
    # -------------------------

    try:

        daily_fine = int(daily_fine)

    except (TypeError, ValueError):

        messages.error(
            request,
            "مبلغ جریمه روزانه باید یک عدد باشد.",
        )

        return redirect("manage_loans")

    if daily_fine < 0:

        messages.error(
            request,
            "مبلغ جریمه روزانه نمی‌تواند منفی باشد.",
        )

        return redirect("manage_loans")

    # -------------------------
    # ثبت امانت
    # -------------------------

    with transaction.atomic():

        book = get_object_or_404(
            Book.objects.select_for_update(),
            pk=book_id,
        )

        user = get_object_or_404(
            User,
            pk=user_id,
        )

        borrowed_count = (
            Loan.objects
            .filter(
                book=book,
                status=Loan.STATUS_BORROWED,
            )
            .count()
        )

        available_copies = (
            book.total_copies -
            borrowed_count
        )

        if available_copies <= 0:

            messages.error(
                request,
                (
                    f"کتاب «{book.title}» "
                    "در حال حاضر هیچ نسخه‌ای برای امانت ندارد."
                ),
            )

            return redirect("manage_loans")

        active_loan_exists = (
            Loan.objects
            .filter(
                user=user,
                book=book,
                status=Loan.STATUS_BORROWED,
            )
            .exists()
        )

        if active_loan_exists:

            messages.error(
                request,
                "این کاربر در حال حاضر همین کتاب را در اختیار دارد.",
            )

            return redirect("manage_loans")

        borrowed_at = timezone.now()

        loan = Loan.objects.create(
            user=user,
            book=book,
            status=Loan.STATUS_BORROWED,
            borrowed_at=borrowed_at,
            due_date=due_date_value,
            daily_fine=daily_fine,
            admin_note=admin_note,
            late_days=0,
            total_fine=0,
        )

        borrowed_count += 1

        book.is_available = (
            borrowed_count <
            book.total_copies
        )

        book.save(
            update_fields=[
                "is_available",
            ]
        )

    # ثبت رویداد در پشتیبانی
    log_activity(
        request=request,
        action=ActivityLog.ACTION_LOAN,
        section="امانت‌ها",
        description=(
            f"کتاب «{book.title}» "
            f"با شماره ثبت امانت "
            f"{loan.registration_number} "
            f"به کاربر "
            f"«{user.get_full_name() or user.username}» "
            "امانت داده شد."
        ),
        object_id=loan.pk,
    )

    messages.success(
        request,
        (
            f"امانت کتاب «{book.title}» "
            f"با شماره ثبت {loan.registration_number} "
            "با موفقیت ثبت شد."
        ),
    )

    return redirect("manage_loans")


@permission_required("blog.manage_loans")
def return_loan(request, pk):

    if request.method != "POST":
        return redirect("manage_loans")

    with transaction.atomic():

        loan = get_object_or_404(
            Loan.objects
            .select_for_update()
            .select_related(
                "book",
                "user",
            ),
            pk=pk,
        )

        if loan.status != Loan.STATUS_BORROWED:

            messages.error(
                request,
                "این امانت قبلاً برگشت داده شده است.",
            )

            return redirect("manage_loans")

        book = (
            Book.objects
            .select_for_update()
            .get(
                pk=loan.book_id,
            )
        )

        now = timezone.now()

        late_days = 0

        if (
            loan.due_date
            and now > loan.due_date
        ):

            difference = (
                now -
                loan.due_date
            )

            late_days = difference.days

        total_fine = (
            late_days *
            loan.daily_fine
        )

        loan.status = (
            Loan.STATUS_AVAILABLE
        )

        loan.returned_at = now
        loan.late_days = late_days
        loan.total_fine = total_fine

        loan.save(
            update_fields=[
                "status",
                "returned_at",
                "late_days",
                "total_fine",
            ]
        )

        borrowed_count = (
            Loan.objects
            .filter(
                book=book,
                status=Loan.STATUS_BORROWED,
            )
            .count()
        )

        book.is_available = (
            borrowed_count <
            book.total_copies
        )

        book.save(
            update_fields=[
                "is_available",
            ]
        )

    # ثبت رویداد در پشتیبانی
    log_activity(
        request=request,
        action=ActivityLog.ACTION_RETURN,
        section="امانت‌ها",
        description=(
            f"کتاب «{loan.book.title}» "
            f"با شماره ثبت امانت "
            f"{loan.registration_number} "
            f"از کاربر "
            f"«{loan.user.get_full_name() or loan.user.username}» "
            f"برگشت داده شد. "
            f"دیرکرد: {late_days} روز، "
            f"مبلغ جریمه: {total_fine:,} تومان."
        ),
        object_id=loan.pk,
    )

    if total_fine > 0:

        messages.success(
            request,
            (
                "بازگشت کتاب با موفقیت ثبت شد. "
                f"تعداد روز دیرکرد: {late_days} روز "
                f"و مبلغ جریمه: {total_fine:,} تومان."
            ),
        )

    else:

        messages.success(
            request,
            (
                "بازگشت کتاب با موفقیت ثبت شد. "
                "این امانت جریمه‌ای نداشت."
            ),
        )

    return redirect("manage_loans")


# =========================================================
# USERS
# =========================================================

@permission_required("accounts.view_user")
def manage_users(request):

    users = (
        User.objects
        .prefetch_related("groups")
        .annotate(
            active_loans_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
            )
        )
        .order_by("-date_joined")
    )

    context = {
        "users": users,
        "add_user_form": AdminUserCreateForm(),
    }

    return render(
        request,
        "dashboard/users.html",
        context
    )



@login_required
def print_user_card(request, pk):
    if not request.user.is_staff and not request.user.is_superuser:
        raise PermissionDenied

    card_user = get_object_or_404(User, pk=pk)

    return render(
        request,
        "accounts/membership_card_selected.html",
        {
            "card_user": card_user,
        },
    )

@permission_required("accounts.change_user")
def edit_user(request, pk):
    user = get_object_or_404(User, pk=pk)

    if user.is_superuser and not request.user.is_superuser:
        messages.error(
            request,
            "شما اجازه ویرایش مدیر سیستم را ندارید."
        )
        return redirect("manage_users")

    can_manage_roles = request.user.has_perm(
        "accounts.manage_roles"
    )

    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        phone_number = normalize_digits(
            request.POST.get("phone_number", "").strip()
        )
        first_name = request.POST.get("first_name", "").strip()
        last_name = request.POST.get("last_name", "").strip()
        father_name = request.POST.get("father_name", "").strip()
        national_code = normalize_digits(
            request.POST.get("national_code", "").strip()
        )
        birth_date = request.POST.get("birth_date", "").strip()
        home_address = request.POST.get("home_address", "").strip()
        password = request.POST.get("password", "")
        is_active = request.POST.get("is_active") == "on"

        def form_error(message):
            messages.error(request, message)
            return redirect("edit_user", pk=user.pk)

        # اعتبارسنجی نام کاربری
        if not username:
            return form_error("نام کاربری نمی‌تواند خالی باشد.")

        if User.objects.filter(username=username).exclude(pk=user.pk).exists():
            return form_error("این نام کاربری قبلاً استفاده شده است.")

        # اعتبارسنجی شماره موبایل
        if phone_number:
            if len(phone_number) != 11 or not phone_number.isdigit():
                return form_error("شماره موبایل باید ۱۱ رقم باشد.")

            if User.objects.filter(
                phone_number=phone_number
            ).exclude(pk=user.pk).exists():
                return form_error(
                    "این شماره موبایل قبلاً ثبت شده است."
                )

        # اعتبارسنجی کد ملی
        if national_code:
            if len(national_code) != 10 or not national_code.isdigit():
                return form_error("کد ملی باید ۱۰ رقم باشد.")

            if User.objects.filter(
                national_code=national_code
            ).exclude(pk=user.pk).exists():
                return form_error(
                    "این کد ملی قبلاً برای کاربر دیگری ثبت شده است."
                )

        # اعتبارسنجی تاریخ تولد
        if birth_date:
            try:
                from datetime import date
                birth_date_value = date.fromisoformat(birth_date)
            except ValueError:
                return form_error(
                    "تاریخ تولد معتبر نیست."
                )
        else:
            birth_date_value = None

        # ذخیره اطلاعات کاربر
        user.username = username
        user.phone_number = phone_number or None
        user.first_name = first_name
        user.last_name = last_name
        user.father_name = father_name or None
        user.national_code = national_code or None
        user.birth_date = birth_date_value
        user.home_address = home_address or None

        if user.pk == request.user.pk or user.is_superuser:
            user.is_active = True
        else:
            user.is_active = is_active

        profile_image = request.FILES.get("profile_image")
        if profile_image:
            user.profile_image = profile_image

        if password:
            user.set_password(password)

        # حفظ نقش‌ها با همان منطق قبلی
        roles_changed = False

        if can_manage_roles and user.pk != request.user.pk:
            selected_roles = request.POST.getlist("roles")

            if (
                "system" in selected_roles
                and not request.user.is_superuser
            ):
                selected_roles.remove("system")
                messages.warning(
                    request,
                    "نقش مدیر سیستم فقط توسط مدیر اصلی قابل اختصاص است."
                )

            set_user_roles(user, selected_roles)
            roles_changed = True

        user.save()

        log_activity(
            request=request,
            action=ActivityLog.ACTION_UPDATE,
            section="کاربران",
            description=(
                f"اطلاعات کاربر «{user.get_full_name() or user.username}» "
                "ویرایش شد."
            ),
            object_id=user.pk,
        )

        if roles_changed:
            log_activity(
                request=request,
                action=ActivityLog.ACTION_ROLE,
                section="نقش‌های کاربران",
                description=(
                    f"نقش‌های کاربر «{user.get_full_name() or user.username}» "
                    "تغییر کرد."
                ),
                object_id=user.pk,
            )

        messages.success(
            request,
            "اطلاعات کاربر با موفقیت ویرایش شد."
        )
        return redirect("manage_users")

    context = {
        "user_obj": user,
        "current_roles": get_user_roles(user),
        "role_choices": ROLE_CHOICES,
        "can_manage_roles": can_manage_roles,
    }

    return render(
        request,
        "dashboard/user_form.html",
        context,
    )


@permission_required("accounts.change_user")
def toggle_user_active(request, pk):
    if request.method != "POST":
        return redirect("manage_users")

    user = get_object_or_404(
        User,
        pk=pk
    )

    if user.pk == request.user.pk:
        messages.error(
            request,
            "نمی‌توانید حساب کاربری خودتان را غیرفعال کنید."
        )

        return redirect("manage_users")

    if user.is_superuser and not request.user.is_superuser:
        messages.error(
            request,
            "شما اجازه غیرفعال کردن مدیر سیستم را ندارید."
        )

        return redirect("manage_users")

    user.is_active = not user.is_active

    user.save(
        update_fields=[
            "is_active",
        ]
    )

    log_activity(
        request=request,
        action=ActivityLog.ACTION_STATUS,
        section="کاربران",
        description=(
            f"حساب کاربر «{user.get_full_name() or user.username}» "
            f"{'فعال' if user.is_active else 'غیرفعال'} شد."
        ),
        object_id=user.pk,
    )

    if user.is_active:
        messages.success(
            request,
            f"حساب «{user.username}» فعال شد."
        )
    else:
        messages.success(
            request,
            f"حساب «{user.username}» غیرفعال شد."
        )

    return redirect("manage_users")


@permission_required("accounts.delete_user")
def delete_user(request, pk):
    if request.method != "POST":
        return redirect("manage_users")

    user = get_object_or_404(
        User,
        pk=pk
    )

    if user.pk == request.user.pk:
        messages.error(
            request,
            "نمی‌توانید حساب کاربری خودتان را حذف کنید."
        )

        return redirect("manage_users")

    if user.is_superuser and not request.user.is_superuser:
        messages.error(
            request,
            "شما اجازه حذف مدیر سیستم را ندارید."
        )

        return redirect("manage_users")

    active_loan_exists = Loan.objects.filter(
        user=user,
        status=Loan.STATUS_BORROWED
    ).exists()

    if active_loan_exists:
        messages.error(
            request,
            "این کاربر هنوز کتاب امانت‌داده‌شده دارد. ابتدا کتاب را پس بگیرید."
        )

        return redirect("manage_users")

    username = user.username
    user_pk = user.pk

    with transaction.atomic():
        user.delete()

    log_activity(
        request=request,
        action=ActivityLog.ACTION_DELETE,
        section="کاربران",
        description=(
            f"کاربر «{username}» حذف شد."
        ),
        object_id=user_pk,
    )

    messages.success(
        request,
        f"کاربر «{username}» با موفقیت حذف شد."
    )

    return redirect("manage_users")

# =========================================================
# SUPPORT / ACTIVITY LOG
# =========================================================

@permission_required("accounts.access_dashboard")
def support_logs(request):

    logs = (
        ActivityLog.objects
        .select_related("actor")
        .all()
    )

    # -----------------------------------------
    # فیلتر نوع عملیات
    # -----------------------------------------

    action = request.GET.get(
        "action",
        ""
    ).strip()

    if action:
        logs = logs.filter(
            action=action
        )

    # -----------------------------------------
    # فیلتر کاربر
    # -----------------------------------------

    actor_id = request.GET.get(
        "actor",
        ""
    ).strip()

    if actor_id:
        logs = logs.filter(
            actor_id=actor_id
        )

    # -----------------------------------------
    # جستجو
    # -----------------------------------------

    query = request.GET.get(
        "q",
        ""
    ).strip()

    if query:
        logs = logs.filter(
            description__icontains=query
        )

    # -----------------------------------------
    # فیلتر زمانی
    # -----------------------------------------

    date_filter = request.GET.get(
        "date",
        "all"
    ).strip()

    now = timezone.localtime(
        timezone.now()
    )

    today = now.date()

    if date_filter == "today":

        logs = logs.filter(
            created_at__date=today
        )

    elif date_filter == "yesterday":

        yesterday = today - timedelta(
            days=1
        )

        logs = logs.filter(
            created_at__date=yesterday
        )

    elif date_filter == "before_yesterday":

        before_yesterday = today - timedelta(
            days=2
        )

        logs = logs.filter(
            created_at__date=before_yesterday
        )

    elif date_filter == "last_7":

        start_date = today - timedelta(
            days=6
        )

        logs = logs.filter(
            created_at__date__gte=start_date
        )

    elif date_filter == "last_30":

        start_date = today - timedelta(
            days=29
        )

        logs = logs.filter(
            created_at__date__gte=start_date
        )

    users = (
        User.objects
        .filter(
            is_active=True
        )
        .order_by(
            "first_name",
            "last_name",
            "username"
        )
    )

    context = {
        "logs": logs,
        "users": users,
        "action_choices": ActivityLog.ACTION_CHOICES,
        "selected_action": action,
        "selected_actor": actor_id,
        "selected_date": date_filter,
        "query": query,
    }

    return render(
        request,
        "dashboard/support.html",
        context
    )
# =========================================================
# DASHBOARD LOGIN
# =========================================================

def dashboard_login(request):
    if request.user.is_authenticated:
        return redirect("dashboard_home")

    if request.method == "GET":
        list(messages.get_messages(request))

    if request.method == "POST":
        username = request.POST.get(
            "username",
            ""
        ).strip()

        password = request.POST.get(
            "password",
            ""
        )

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is None:
            messages.error(
                request,
                "نام کاربری یا رمز عبور اشتباه است."
            )

            return render(
                request,
                "dashboard/login.html"
            )

        if not user.is_staff:
            messages.error(
                request,
                "شما اجازه دسترسی به پنل مدیریت را ندارید."
            )

            return render(
                request,
                "dashboard/login.html"
            )

        login(request, user)

        log_activity(
            request=request,
            action=ActivityLog.ACTION_LOGIN,
            section="ورود به پنل",
            description=(
                f"کاربر «{user.get_full_name() or user.username}» "
                f"وارد پنل مدیریت شد."
            ),
            object_id=user.pk,
        )

        return redirect("dashboard_home")

    return render(
        request,
        "dashboard/login.html"
    )





@permission_required("accounts.access_dashboard")
def export_support_excel(request):

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from django.http import HttpResponse
    import jdatetime

    logs = (
        ActivityLog.objects
        .select_related("actor")
        .order_by("-created_at")
    )

    action = request.GET.get("action", "").strip()

    if action:
        logs = logs.filter(
            action=action
        )

    actor_id = request.GET.get("actor", "").strip()

    if actor_id:
        logs = logs.filter(
            actor_id=actor_id
        )

    query = request.GET.get("q", "").strip()

    if query:
        logs = logs.filter(
            description__icontains=query
        )

    date_filter = request.GET.get(
        "date",
        "all"
    ).strip()

    now = timezone.localtime(
        timezone.now()
    )

    today = now.date()

    if date_filter == "today":

        logs = logs.filter(
            created_at__date=today
        )

    elif date_filter == "yesterday":

        yesterday = (
            today -
            timedelta(days=1)
        )

        logs = logs.filter(
            created_at__date=yesterday
        )

    elif date_filter == "before_yesterday":

        before_yesterday = (
            today -
            timedelta(days=2)
        )

        logs = logs.filter(
            created_at__date=before_yesterday
        )

    elif date_filter == "last_7":

        start_date = (
            today -
            timedelta(days=6)
        )

        logs = logs.filter(
            created_at__date__gte=start_date
        )

    elif date_filter == "last_30":

        start_date = (
            today -
            timedelta(days=29)
        )

        logs = logs.filter(
            created_at__date__gte=start_date
        )


    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "گزارش فعالیت‌ها"

    worksheet.sheet_view.rightToLeft = True


    headers = [
        "تاریخ و ساعت",
        "انجام‌دهنده",
        "عملیات",
        "بخش",
        "توضیحات",
        "شناسه رکورد",
        "آدرس IP",
    ]


    worksheet.append(headers)


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    for log in logs:

        actor_name = "سیستم"

        if log.actor:

            if log.actor.get_full_name():

                actor_name = (
                    log.actor.get_full_name()
                )

            else:

                actor_name = (
                    log.actor.username
                )


        if log.created_at:

            created_at = (
                timezone.localtime(
                    log.created_at
                )
            )

            jalali_datetime = (
                jdatetime.datetime.fromgregorian(
                    datetime=created_at
                )
            )

            date_text = (
                jalali_datetime.strftime(
                    "%Y/%m/%d %H:%M:%S"
                )
            )

        else:

            date_text = ""


        worksheet.append(
            [
                date_text,
                actor_name,
                log.get_action_display(),
                log.section or "-",
                log.description,
                log.object_id or "",
                log.ip_address or "",
            ]
        )


    column_widths = {
        "A": 22,
        "B": 25,
        "C": 18,
        "D": 22,
        "E": 60,
        "F": 18,
        "G": 20,
    }


    for column, width in column_widths.items():

        worksheet.column_dimensions[
            column
        ].width = width


    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="support_logs.xlsx"'
    )


    workbook.save(
        response
    )


    return response



@permission_required("accounts.change_user")
def add_user(request):

    if request.method != "POST":
        return redirect("manage_users")

    form = AdminUserCreateForm(
        request.POST,
        request.FILES
    )

    if form.is_valid():

        user = form.save(
            commit=False
        )

        user.phone_number = (
            form.cleaned_data["phone_number"]
        )

        user.is_active = True
        user.is_staff = False
        user.is_superuser = False

        user.save()

        log_activity(
            request=request,
            action=ActivityLog.ACTION_CREATE,
            section="کاربران",
            description=(
                f"کاربر «{user.username}» "
                f"با موفقیت توسط مدیر ایجاد شد."
            ),
            object_id=user.pk,
        )

        messages.success(
            request,
            f"کاربر «{user.username}» با موفقیت اضافه شد."
        )

        return redirect(
            "manage_users"
        )

    # اگر فرم خطا داشت، دوباره صفحه مدیریت کاربران
    # را همراه با همان خطاها نمایش می‌دهیم.

    users = (
        User.objects
        .prefetch_related("groups")
        .annotate(
            active_loans_count=Count(
                "loans",
                filter=Q(
                    loans__status=Loan.STATUS_BORROWED
                ),
            )
        )
        .order_by("-date_joined")
    )

    context = {
        "users": users,
        "add_user_form": form,
    }

    return render(
        request,
        "dashboard/users.html",
        context
    )